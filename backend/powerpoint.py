"""The PowerPoint export: a deck as a `.pptx` whose text can be edited, built from the page the browser rendered.

Each slide is a picture of the page with its text made invisible, as the slide's background, with every block of
text laid over it as a text box at the position, size, font and colour the browser gave it. Both halves come
from the runtime's scene mode (`frontend/src/scene.ts`): the page opened with `?scene` measures its text into a
JSON scene the chrome service reads out of the DOM (`POST /scene/`), and a screenshot of the same URL, one per
page at a window of exactly the slide's size, is the page without its glyphs. The decorations CSS gives the
text (the markdown themes' heading prefixes and bullets, the backgrounds of inline code, borders, gradients,
SVG components, images) stay in the picture; the words become text boxes, so retyping a slide is a matter of
clicking it. Moving a text box leaves its background where it was: that is the trade the hybrid makes.

Positions are CSS pixels of the unscaled slide, which the scene also gives the size of; a slide of the same size
in inches (96 px to the inch) keeps every font at its CSS size in points. Decks only: a document's or page
artifact's pages are not slides.
"""

from __future__ import annotations

import asyncio
import io
from collections.abc import Sequence
from typing import Literal

import render
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.slide import Slide
from pptx.util import Emu, Pt
from pydantic import BaseModel

import workspace

# A CSS pixel is 1/96 inch; an EMU is 1/914400 inch; a point is 1/72 inch.
EMU_PER_PX = 9525
PT_PER_PX = 0.75
# The device pixel ratio of the slide pictures: twice the CSS size, so they stay sharp on a projector.
PICTURE_SCALE = 2
# How many pages are photographed at once; each is a Chrome run on the chrome service.
CONCURRENCY = 4
# PowerPoint wraps with its own fonts and metrics: a little width keeps a line that just fit from breaking.
WRAP_SLACK = 0.03
# What a CSS generic family becomes; the browser's first named family is used as it is, PowerPoint substituting
# whatever it lacks, except for the macOS monospace names, which Office does not have and would replace with a
# proportional face.
GENERIC_FONTS = {
    'sans-serif': 'Calibri',
    'serif': 'Georgia',
    'monospace': 'Consolas',
    'cursive': 'Comic Sans MS',
    'fantasy': 'Impact',
    'system-ui': 'Calibri',
}
FONT_SUBSTITUTES = {'SFMono-Regular': 'Consolas', 'SF Mono': 'Consolas', 'Menlo': 'Consolas', 'Monaco': 'Consolas'}
ALIGNMENTS = {
    'left': PP_ALIGN.LEFT,
    'center': PP_ALIGN.CENTER,
    'right': PP_ALIGN.RIGHT,
    'justify': PP_ALIGN.JUSTIFY,
}
MEDIA_TYPE = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'


class Run(BaseModel):
    """A stretch of text in one style; `text` of a newline alone is a line break. Mirrors `SceneRun` in scene.ts."""

    text: str
    font: str
    size: float
    bold: bool
    italic: bool
    underline: bool
    strike: bool
    color: str
    spacing: float
    href: str | None = None


class Block(BaseModel):
    """One text box: its position and size in slide pixels and its runs. Mirrors `SceneBlock` in scene.ts."""

    x: float
    y: float
    w: float
    h: float
    align: Literal['left', 'center', 'right', 'justify']
    lineHeight: float
    runs: list[Run]


class Page(BaseModel):
    index: int
    blocks: list[Block]


class Scene(BaseModel):
    """What the runtime measured: the slide size in CSS pixels and the text blocks of every page."""

    width: float
    height: float
    pages: list[Page]


async def export(found: workspace.Artifact) -> bytes:
    """The deck as a `.pptx`: its scene and one picture per page from the chrome service, assembled in a thread.

    The page is built first, so a broken artifact fails as a build error; the chrome service's failures are
    `render.RenderError`s.
    """
    async with workspace.open_artifact(found) as directory:
        await render.ensure_built(found, directory)
    scene = Scene.model_validate_json(await render.call_chrome('/scene/', {'url': render.print_url(found, scene=True)}))
    width, height = round(scene.width), round(scene.height)
    limit = asyncio.Semaphore(CONCURRENCY)

    async def picture(page: Page) -> bytes:
        async with limit:
            body: dict[str, object] = {
                'url': render.print_url(found, page.index, scene=True),
                'width': width,
                'height': height,
                'scale': PICTURE_SCALE,
            }
            return await render.call_chrome('/screenshot/', body)

    pictures = await asyncio.gather(*(picture(page) for page in scene.pages))
    return await asyncio.to_thread(build_pptx, scene, pictures, found.title)


def build_pptx(scene: Scene, pictures: Sequence[bytes], title: str = '') -> bytes:
    """Assemble the presentation: one slide per page of `scene`, each its picture behind its text boxes."""
    prs = Presentation()
    prs.slide_width = Emu(round(scene.width * EMU_PER_PX))
    prs.slide_height = Emu(round(scene.height * EMU_PER_PX))
    prs.core_properties.title = title
    blank = prs.slide_layouts[6]
    for page, picture in zip(scene.pages, pictures, strict=True):
        slide = prs.slides.add_slide(blank)
        set_background(slide, picture)
        for n, block in enumerate(page.blocks, 1):
            add_block(slide, block, n)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def set_background(slide: Slide, png: bytes) -> None:
    """Make `png` the slide's background fill, stretched to the slide: part of the slide, not a shape to drag.

    python-pptx has no API for a picture background, so the `<p:bg>` element is written by hand, first in the
    slide's common data as the schema orders it.
    """
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(png))
    # python-pptx's XML layer is untyped from here down.
    background = parse_xml(  # pyright: ignore[reportUnknownVariableType]
        f'<p:bg {nsdecls("p", "a", "r")}><p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/>'
        '<a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>'
    )
    slide.element.cSld.insert(0, background)  # pyright: ignore[reportUnknownMemberType]


def add_block(slide: Slide, block: Block, n: int) -> None:
    """Add `block` as a text box: its runs in one paragraph, line breaks where the browser broke lines."""
    slack = block.w * WRAP_SLACK
    # Extra width to the right of left-aligned text moves nothing; centred text must keep its centre.
    x = block.x - (slack if block.align == 'right' else slack / 2 if block.align == 'center' else 0)
    shape = slide.shapes.add_textbox(emu(x), emu(block.y), emu(block.w + slack), emu(block.h))
    shape.name = f'Text {n}'
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Emu(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = ALIGNMENTS[block.align]
    paragraph.line_spacing = Pt(block.lineHeight * PT_PER_PX)
    paragraph.space_before = paragraph.space_after = Pt(0)
    for run in block.runs:
        if run.text == '\n':
            paragraph.add_line_break()
            continue
        added = paragraph.add_run()
        added.text = run.text
        font = added.font
        font.name = font_name(run.font)
        font.size = Pt(run.size * PT_PER_PX)
        font.bold = run.bold
        font.italic = run.italic
        font.underline = run.underline
        font.color.rgb = rgb(run.color)
        # Strike-through and character spacing have no python-pptx API: they are attributes of the run properties.
        properties = added._r.get_or_add_rPr()  # pyright: ignore[reportPrivateUsage]
        if run.strike:
            properties.set('strike', 'sngStrike')  # pyright: ignore[reportUnknownMemberType]
        if run.spacing:
            # Character spacing is in hundredths of a point.
            properties.set('spc', str(round(run.spacing * PT_PER_PX * 100)))  # pyright: ignore[reportUnknownMemberType]
        if run.href:
            added.hyperlink.address = run.href


def font_name(family: str) -> str:
    """The PowerPoint font for a CSS family: generics and the macOS monospace names mapped, anything else as is."""
    return GENERIC_FONTS.get(family) or FONT_SUBSTITUTES.get(family) or family


def rgb(color: str) -> RGBColor:
    """`#rrggbb` as a python-pptx colour."""
    value = int(color.lstrip('#'), 16)
    return RGBColor(value >> 16 & 0xFF, value >> 8 & 0xFF, value & 0xFF)


def emu(px: float) -> Emu:
    """CSS pixels of the slide to EMUs."""
    return Emu(round(px * EMU_PER_PX))
