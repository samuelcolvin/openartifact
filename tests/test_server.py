"""Tests for `backend/server.py`: the HTTP routes with the whole app running, and MCP over real HTTP."""

from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import threading
import time
import uuid
import zipfile
from collections.abc import Awaitable, Callable, Iterator
from pathlib import Path
from typing import TypeVar

import httpx2
import logfire
import powerpoint
import pytest
import render
import uvicorn
import word
from conftest import DEV_TOKEN
from docx import Document
from fastapi import FastAPI, Request
from fastapi.responses import Response
from fastapi.testclient import TestClient
from fastmcp import Client
from logfire.testing import CaptureLogfire
from PIL import Image
from pptx import Presentation

import auth
import db
import login
import server
import upload
import workspace

ROOT = Path(__file__).resolve().parent.parent
STARTER = ROOT / 'examples' / 'starter'
T = TypeVar('T')


@pytest.fixture
def client(server_env: None) -> Iterator[TestClient]:
    with TestClient(server.app) as client:
        yield client


async def seed_starter() -> tuple[auth.Principal, workspace.Artifact]:
    """Import the starter deck as the seed user's (private) artifact; runs in the app's loop."""
    async with db.pool().acquire() as conn, conn.transaction():
        principal = await auth.upsert_user(conn, sub='seed', email='seed@example.com', name='Seed', picture=None)
    return principal, await workspace.import_directory(principal.workspace_id, 'Starter', 'deck', STARTER)


def sign_in(client: TestClient, principal: auth.Principal) -> None:
    """Give the client the session cookie a browser gets from `/login`."""
    client.cookies.set(login.SESSION_COOKIE, login.make_session(principal.user_id, FAR_FUTURE))


def sign_out(client: TestClient) -> None:
    client.cookies.delete(login.SESSION_COOKIE)


def in_app(client: TestClient, fn: Callable[[], Awaitable[T]]) -> T:
    """Run a coroutine in the app's loop, so it uses the app's pool, store and checkout cache."""
    assert client.portal is not None
    return client.portal.call(fn)


def starter(client: TestClient) -> workspace.Artifact:
    """Seed the starter deck through the running app, signed in as its owner (the artifact is private)."""
    principal, artifact = in_app(client, seed_starter)
    sign_in(client, principal)
    return artifact


def test_index(client: TestClient):
    assert client.get('/api/').json() == {'mcp': '/mcp/', 'runtime': '/openartifact.js', 'login': '/login', 'app': '/'}


def test_health(client: TestClient):
    assert client.get('/health/').json() == {'status': 'ok'}


def test_runtime_js(client: TestClient):
    response = client.get('/openartifact.js')
    assert response.status_code == 200
    assert response.headers['content-type'].startswith('text/javascript')
    assert response.content == server.RUNTIME_JS_PATH.read_bytes()


def test_favicon(client: TestClient):
    """The platform's mark at the root, under both names a browser asks for."""
    for path in '/favicon.svg', '/favicon.ico':
        response = client.get(path)
        assert response.status_code == 200
        assert response.headers['content-type'].startswith('image/svg+xml')
        assert response.content == server.FAVICON_PATH.read_bytes()


def test_unknown_artifacts_are_404(client: TestClient):
    assert client.get(f'/artifacts/{uuid.uuid4()}/').status_code == 404
    assert client.get(f'/artifacts/{uuid.uuid4()}/assets/logo.svg').status_code == 404
    assert client.get('/artifacts/Bad%20Name/').status_code == 404
    assert client.get('/artifacts/not-a-uuid/assets/logo.svg').status_code == 404


def test_artifact_page_is_built_on_demand(client: TestClient):
    artifact = starter(client)
    directory = workspace.checkout_path(artifact.id)
    assert not (directory / 'dist').exists()
    page = client.get(f'/artifacts/{artifact.id}/')
    assert page.status_code == 200
    assert page.headers['content-type'].startswith('text/html')
    assert '<script src="/openartifact.js"></script>' in page.text
    assert 'data:' not in page.text
    assert (directory / 'dist' / 'index.html').is_file()
    # A later request serves the built file; deleting it triggers another build.
    shutil.rmtree(directory / 'dist')
    assert client.get(f'/artifacts/{artifact.id}/').status_code == 200


def test_artifact_build_failure_is_422(client: TestClient):
    artifact = starter(client)
    directory = workspace.checkout_path(artifact.id)

    async def break_it() -> None:
        async with workspace.edit(artifact.id, 'break') as tx:
            (tx.path / 'main.md').write_text('# hi\n\n<component src="Nope.html"></component>\n')

    in_app(client, break_it)
    response = client.get(f'/artifacts/{artifact.id}/')
    assert response.status_code == 422
    assert 'component not found' in response.json()['detail']
    assert not (directory / 'dist').exists()


def test_artifact_images_are_served(client: TestClient):
    artifact = starter(client)
    logo = client.get(f'/artifacts/{artifact.id}/assets/logo.svg')
    assert logo.status_code == 200
    assert logo.headers['content-type'].startswith('image/svg+xml')
    assert logo.content == (STARTER / 'assets' / 'logo.svg').read_bytes()
    # Fonts declared in styles.css are served too (any file in the directory with an allowed extension).
    directory = workspace.checkout_path(artifact.id)
    (directory / 'assets' / 'body.woff2').write_bytes(b'wOF2')
    assert client.get(f'/artifacts/{artifact.id}/assets/body.woff2').status_code == 200


def test_artifact_sources_are_served_as_text(client: TestClient):
    artifact = starter(client)
    for path, media_type in (
        ('main.md', 'text/markdown; charset=utf-8'),
        ('artifact.toml', 'text/plain; charset=utf-8'),
        ('styles.css', 'text/plain; charset=utf-8'),
        ('components/Hero.html', 'text/plain; charset=utf-8'),
    ):
        response = client.get(f'/artifacts/{artifact.id}/{path}')
        assert response.status_code == 200, path
        assert response.headers['content-type'] == media_type, path
        assert response.content == (STARTER / path).read_bytes(), path
    # The page advertises the markdown export, with a note telling agents to read it rather than the HTML.
    page = client.get(f'/artifacts/{artifact.id}/')
    assert (
        f'<link rel="alternate" type="text/markdown" href="../{artifact.id}.md" title="Markdown source">' in page.text
    )
    assert 'do not parse this HTML' in page.text and f'\n      ../{artifact.id}.md\n' in page.text
    assert 'id="artifact-markdown"' in page.text
    # The build output is not a source file: the page is served at the directory URL only.
    for path in ('dist/index.html', 'assets/../dist/index.html'):
        assert client.get(f'/artifacts/{artifact.id}/{path}').status_code == 404, path


def test_artifact_media_cannot_reach_other_artifacts(client: TestClient):
    artifact = starter(client)
    other = starter(client)
    # `..` that leaves the artifact directory is refused, even towards another artifact's image. The segments
    # are percent-encoded because the HTTP client collapses a literal `..` before sending.
    assert client.get(f'/artifacts/{artifact.id}/%2e%2e/{other.id}/assets/logo.svg').status_code == 404
    assert client.get(f'/artifacts/{artifact.id}/assets/%2e%2e/%2e%2e/{other.id}/assets/logo.svg').status_code == 404
    # ...while `..` that stays inside it is just a path to the same file.
    assert client.get(f'/artifacts/{artifact.id}/assets/../assets/logo.svg').status_code == 200


def test_artifact_media_cannot_escape_via_symlink(client: TestClient, tmp_path: Path):
    artifact = starter(client)
    outside = tmp_path / 'outside.png'
    outside.write_bytes(b'png')
    directory = workspace.checkout_path(artifact.id)
    (directory / 'assets' / 'link.png').symlink_to(outside)
    assert client.get(f'/artifacts/{artifact.id}/assets/link.png').status_code == 404


def test_artifact_redirects_to_trailing_slash(client: TestClient):
    response = client.get(f'/artifacts/{uuid.uuid4()}', follow_redirects=False)
    assert response.status_code == 307


def starter_sources() -> list[str]:
    """The starter deck's source files, relative, as the server lists them."""
    return sorted(
        p.relative_to(STARTER).as_posix() for p in STARTER.rglob('*') if p.is_file() and 'dist' not in p.parts
    )


def test_artifact_markdown_has_a_frontmatter_summary(client: TestClient):
    artifact = starter(client)
    response = client.get(f'/artifacts/{artifact.id}.md')
    assert response.status_code == 200
    assert response.headers['content-type'] == 'text/markdown; charset=utf-8'
    empty, front, body = response.text.split('---\n', 2)
    assert empty == ''
    lines = front.splitlines()
    # Title and theme come from artifact.toml, which outranks the row's title ("Starter").
    assert lines[:4] == [
        f'id: "{artifact.id}"',
        'title: "OpenArtifact Starter"',
        'type: "deck"',
        'theme: "markdown-dark"',
    ]
    assert f'url: "http://127.0.0.1:8765/artifacts/{artifact.id}/"' in lines
    assert any(line.startswith('created_at: "') for line in lines)
    files = [line.removeprefix('  - ').strip('"') for line in lines[lines.index('files:') + 1 :]]
    assert files == starter_sources()
    assert body == (STARTER / 'main.md').read_text(encoding='utf-8')
    # The summary does not need the artifact to build: a broken artifact.toml falls back to the row.
    directory = workspace.checkout_path(artifact.id)
    (directory / 'artifact.toml').write_text('this is not toml')
    response = client.get(f'/artifacts/{artifact.id}.md')
    assert response.status_code == 200
    assert response.text.splitlines()[1:4] == [f'id: "{artifact.id}"', 'title: "Starter"', 'type: "deck"']


def last_commit(client: TestClient, artifact: workspace.Artifact) -> str:
    """Short sha of the last commit that touched the artifact, as the download names use it."""
    return in_app(client, lambda: workspace.artifact_git(artifact.id, 'rev-parse', 'HEAD')).strip()[:7]


def test_page_url_serves_markdown_to_clients_that_prefer_text(client: TestClient):
    artifact = starter(client)
    url = f'/artifacts/{artifact.id}/'
    markdown = client.get(f'{url[:-1]}.md').text
    for accept in ('text/markdown', 'text/plain', 'text/plain;q=0.9, text/html;q=0.8', 'text/markdown, */*;q=0.1'):
        response = client.get(url, headers={'accept': accept})
        assert response.headers['content-type'] == 'text/markdown; charset=utf-8', accept
        assert response.headers['vary'] == 'Accept, Cookie'
        assert response.text == markdown
    # Browsers, `*/*`, `text/*` and no header at all get the page: HTML wins ties.
    browser = 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8'
    for headers in ({'accept': browser}, {'accept': '*/*'}, {'accept': 'text/*'}, {}):
        response = client.get(url, headers=headers)
        assert response.headers['content-type'].startswith('text/html'), headers
        assert response.headers['vary'] == 'Accept, Cookie'
        assert 'id="artifact-markdown"' in response.text


def test_prefers_text():
    assert server.prefers_text('text/markdown')
    assert server.prefers_text('text/plain, text/html;q=0.5')
    assert not server.prefers_text(None)
    assert not server.prefers_text('')
    assert not server.prefers_text('text/html, text/plain')
    assert not server.prefers_text('text/plain;q=0, text/html;q=0')
    assert not server.prefers_text('text/plain;q=nonsense')
    assert server.prefers_text('TEXT/PLAIN ; q=0.5')


def test_artifact_zip_is_a_repository_with_the_sources_and_history(client: TestClient, tmp_path: Path):
    artifact = starter(client)
    # A second commit, through the upload route, so there is a history to carry.
    assert client.put(upload_to(artifact, 'main.md', 6), content=b'# new\n').status_code == 200
    # Build first, so there is a dist/ to leave out.
    assert client.get(f'/artifacts/{artifact.id}/').status_code == 200
    response = client.get(f'/artifacts/{artifact.id}.zip')
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/zip'
    # Named after the artifact.toml title and the commit it came from.
    sha = last_commit(client, artifact)
    assert response.headers['content-disposition'] == (
        f'attachment; filename="OpenArtifact Starter {sha}.zip"; filename*=UTF-8\'\'OpenArtifact%20Starter%20{sha}.zip'
    )
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        names = archive.namelist()
        sources = [n for n in names if not n.startswith(f'{artifact.id}/.git/')]
        assert sources == [f'{artifact.id}/{path}' for path in starter_sources()]
        assert archive.read(f'{artifact.id}/main.md') == b'# new\n'
        assert f'{artifact.id}/.git/HEAD' in names
        archive.extractall(tmp_path)
    # The folder is a working clone of the artifact alone: its history, with the original messages and dates,
    # and the artifact's files at the root, with nothing to commit.
    repo = tmp_path / str(artifact.id)
    log = in_app(client, lambda: workspace.git('log', '--format=%s', cwd=repo))
    assert log.splitlines() == ['upload: main.md', f'import: {artifact.id}']
    assert in_app(client, lambda: workspace.git('status', '--porcelain', cwd=repo)) == ''
    assert (
        in_app(client, lambda: workspace.git('show', 'HEAD~1:main.md', cwd=repo)) == (STARTER / 'main.md').read_text()
    )
    assert in_app(client, lambda: workspace.git('ls-tree', '--name-only', 'HEAD', cwd=repo)).split() == sorted(
        {p.split('/')[0] for p in starter_sources()}
    )
    # A clone with no remote, and dist/ excluded locally as in the checkout.
    assert in_app(client, lambda: workspace.git('remote', cwd=repo)) == ''
    assert (repo / '.git' / 'info' / 'exclude').read_text() == 'dist/\n'


def test_download_names_survive_awkward_titles(client: TestClient):
    artifact = starter(client)
    directory = workspace.checkout_path(artifact.id)
    (directory / 'artifact.toml').write_text('title = "Zoë \\"Q\\" / 2026 <v1>"\n')
    sha = last_commit(client, artifact)
    disposition = client.get(f'/artifacts/{artifact.id}.zip').headers['content-disposition']
    # Quotes, slashes and brackets become spaces; the ASCII fallback drops the diaeresis, the UTF-8 name keeps it.
    assert disposition == (
        f'attachment; filename="Zo Q 2026 v1 {sha}.zip"; filename*=UTF-8\'\'Zo%C3%AB%20Q%202026%20v1%20{sha}.zip'
    )
    # No title anywhere usable: the row's title, then the id.
    (directory / 'artifact.toml').write_text('title = ""\n')
    assert f'filename="Starter {sha}.zip"' in client.get(f'/artifacts/{artifact.id}.zip').headers['content-disposition']


def test_artifact_pdf_needs_the_chrome_service(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    artifact = starter(client)
    monkeypatch.delenv('OPENARTIFACT_CHROME_URL', raising=False)
    response = client.get(f'/artifacts/{artifact.id}.pdf')
    assert response.status_code == 503
    assert 'OPENARTIFACT_CHROME_URL' in response.json()['detail']


def stub_chrome(monkeypatch: pytest.MonkeyPatch, status: int, body: bytes, scene: str = '') -> list[str]:
    """Point the server at an in-process stand-in for the chrome service; returns the URLs it was asked to print.

    `scene` is what its `/scene/` endpoint answers, for the PowerPoint export.
    """
    asked: list[str] = []
    stub = FastAPI()

    @stub.post('/scene/')
    async def read_scene(request: Request) -> Response:
        asked.append((await request.json())['url'])
        return Response(scene, status_code=status, media_type='application/json' if status == 200 else 'text/plain')

    @stub.post('/pdf/')
    async def print_pdf(request: Request) -> Response:
        asked.append((await request.json())['url'])
        return Response(body, status_code=status, media_type='application/pdf' if status == 200 else 'text/plain')

    @stub.post('/screenshot/')
    async def take_screenshot(request: Request) -> Response:
        sent = await request.json()
        asked.append(f'{sent["url"]} {sent["width"]}x{sent["height"]}')
        return Response(body, status_code=status, media_type='image/png' if status == 200 else 'text/plain')

    monkeypatch.setenv('OPENARTIFACT_CHROME_URL', 'http://chrome:8766/')
    monkeypatch.setenv('OPENARTIFACT_INTERNAL_URL', 'http://app:8765')
    monkeypatch.setattr(render, 'chrome_client', lambda: httpx2.AsyncClient(transport=httpx2.ASGITransport(app=stub)))
    return asked


def test_artifact_pdf_is_printed_by_the_chrome_service(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    artifact = starter(client)
    asked = stub_chrome(monkeypatch, 200, b'%PDF-1.4 stub')
    response = client.get(f'/artifacts/{artifact.id}.pdf')
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == 'application/pdf'
    sha = last_commit(client, artifact)
    assert response.headers['content-disposition'].startswith(f'attachment; filename="OpenArtifact Starter {sha}.pdf"')
    assert response.content == b'%PDF-1.4 stub'
    # The chrome service was given the internal address of the print pass, and the page had been built for it.
    [page_url] = asked
    match = re.fullmatch(rf'http://app:8765/print/([^/]+)/artifacts/{artifact.id}/', page_url)
    assert match, page_url
    token = match.group(1)
    directory = workspace.checkout_path(artifact.id)
    assert (directory / 'dist' / 'index.html').is_file()
    # The pass lets Chrome, which has no session, fetch the private page and its media; nothing else does.
    sign_out(client)
    assert client.get(f'/print/{token}/artifacts/{artifact.id}/').status_code == 200
    assert client.get(f'/print/{token}/artifacts/{artifact.id}/assets/logo.svg').status_code == 200
    assert client.get(f'/print/{token}x/artifacts/{artifact.id}/').status_code == 403
    other = in_app(client, seed_starter)[1]
    assert client.get(f'/print/{token}/artifacts/{other.id}/').status_code == 403
    assert client.get(f'/artifacts/{artifact.id}/', headers={'accept': 'application/json'}).status_code == 401


def test_artifact_pdf_reports_chrome_failures(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    artifact = starter(client)
    stub_chrome(monkeypatch, 502, b'Chrome exited with code 3')
    response = client.get(f'/artifacts/{artifact.id}.pdf')
    assert response.status_code == 502
    assert response.json()['detail'] == 'chrome service failed (502): Chrome exited with code 3'
    # A page that does not build is reported before anything is sent to the chrome service.
    asked = stub_chrome(monkeypatch, 200, b'%PDF-1.4 stub')
    directory = workspace.checkout_path(artifact.id)
    shutil.rmtree(directory / 'dist', ignore_errors=True)
    (directory / 'main.md').write_text('# a\n---\n# b\n')
    response = client.get(f'/artifacts/{artifact.id}.pdf')
    assert response.status_code == 422
    assert asked == []


# The stub stands in for the chrome service, which accepts incoming trace context (`distributed_tracing=True` in
# `chrome/main.py`); here it runs under the test's Logfire configuration, which warns about it instead.
@pytest.mark.filterwarnings('ignore:Found propagated trace context')
def test_artifact_pdf_request_carries_the_trace(
    capfire: CaptureLogfire, client: TestClient, monkeypatch: pytest.MonkeyPatch
):
    """With httpx instrumented (as `main.py` does), the chrome service is called inside the request's trace."""
    artifact = starter(client)
    seen: list[dict[str, str]] = []
    stub = FastAPI()

    @stub.post('/pdf/')
    async def print_pdf(request: Request) -> Response:
        seen.append(dict(request.headers))
        return Response(b'%PDF-1.4 stub', media_type='application/pdf')

    def instrumented() -> httpx2.AsyncClient:
        http = httpx2.AsyncClient(transport=httpx2.ASGITransport(app=stub))
        # Logfire's signature mentions `httpx.Client`, and only httpx2 is installed here, so pyright sees Unknown.
        logfire.instrument_httpx(http)  # pyright: ignore[reportUnknownMemberType]
        return http

    monkeypatch.setenv('OPENARTIFACT_CHROME_URL', 'http://chrome:8766')
    monkeypatch.setattr(render, 'chrome_client', instrumented)
    assert client.get(f'/artifacts/{artifact.id}.pdf').status_code == 200
    [headers] = seen
    assert 'traceparent' in headers
    # One trace: the request span, the httpx client span inside it (named by method alone, as the stable HTTP
    # conventions say for a client; FastAPI's `fastapi.endpoint` span sits between them), and the stub's own
    # server span as the client span's child, since the header reached it.
    spans = {span['name']: span for span in capfire.exporter.exported_spans_as_dict()}
    request_span = spans['GET /artifacts/{artifact_id}.pdf']
    client_span = spans['POST']
    stub_span = spans['POST /pdf/']
    trace_id = request_span['context']['trace_id']
    assert client_span['context']['trace_id'] == trace_id
    # The stub's parent came in over the wire, so it is the client span's context marked remote.
    assert stub_span['parent'] == {**client_span['context'], 'is_remote': True}
    assert headers['traceparent'].split('-')[1] == format(trace_id, '032x')


def test_artifact_png_is_one_page_by_the_chrome_service(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    artifact = starter(client)
    asked = stub_chrome(monkeypatch, 200, b'\x89PNG stub')
    response = client.get(f'/artifacts/{artifact.id}.png?page=3')
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == 'image/png' and response.content == b'\x89PNG stub'
    assert 'content-disposition' not in response.headers  # shown inline
    assert response.headers['cache-control'] == 'private'
    # The chrome service got the print pass with the page as the hash, at the deck's window size.
    [page_url] = asked
    assert re.fullmatch(rf'http://app:8765/print/[^/]+/artifacts/{artifact.id}/#3 1600x900', page_url)
    # Page 1 by default; pages are 1-based.
    assert client.get(f'/artifacts/{artifact.id}.png').status_code == 200
    assert asked[-1].endswith('/#1 1600x900')
    assert client.get(f'/artifacts/{artifact.id}.png?page=0').status_code == 422
    # The same access rules as the page: the starter is private, so a colleague is refused and a visitor sent away.
    colleague, _ = org_members(client)
    sign_in(client, colleague)
    assert client.get(f'/artifacts/{artifact.id}.png').status_code == 403
    sign_out(client)
    assert client.get(f'/artifacts/{artifact.id}.png').status_code == 401
    # No chrome service, no image.
    sign_in(client, in_app(client, seed_starter)[0])  # upserts the same seed user again
    monkeypatch.delenv('OPENARTIFACT_CHROME_URL')
    assert client.get(f'/artifacts/{artifact.id}.png').status_code == 503


SCENE = {
    'width': 1055.9,
    'height': 594.1,
    'pages': [
        {
            'index': n,
            'blocks': [
                {
                    'x': 81,
                    'y': 83,
                    'w': 929,
                    'h': 52,
                    'align': 'left',
                    'lineHeight': 51.52,
                    'runs': [
                        {
                            'text': f'Page {n}',
                            'font': 'Segoe UI',
                            'size': 44.8,
                            'bold': True,
                            'italic': False,
                            'underline': False,
                            'strike': False,
                            'color': '#fbffea',
                            'spacing': 0,
                        }
                    ],
                }
            ],
        }
        for n in (1, 2)
    ],
}


def test_artifact_pptx_is_assembled_from_the_chrome_services_scene_and_pictures(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
):
    artifact = starter(client)
    picture = io.BytesIO()
    Image.new('RGB', (4, 2)).save(picture, 'PNG')
    asked = stub_chrome(monkeypatch, 200, picture.getvalue(), json.dumps(SCENE))
    response = client.get(f'/artifacts/{artifact.id}.pptx')
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == powerpoint.MEDIA_TYPE
    sha = last_commit(client, artifact)
    assert response.headers['content-disposition'].startswith(f'attachment; filename="OpenArtifact Starter {sha}.pptx"')
    assert response.headers['cache-control'] == 'private'
    # The scene was read from the print pass in scene mode, then each of its pages photographed at the slide's size.
    scene_url, *pictures = asked
    assert re.fullmatch(rf'http://app:8765/print/[^/]+/artifacts/{artifact.id}/\?scene', scene_url)
    assert sorted(pictures) == [f'{scene_url}#1 1056x594', f'{scene_url}#2 1056x594']
    prs = Presentation(io.BytesIO(response.content))
    texts = [shape.text_frame.text for slide in prs.slides for shape in slide.shapes]  # pyright: ignore
    assert texts == ['Page 1', 'Page 2']
    # The same access rules as the page.
    colleague, _ = org_members(client)
    sign_in(client, colleague)
    assert client.get(f'/artifacts/{artifact.id}.pptx').status_code == 403
    sign_out(client)
    assert client.get(f'/artifacts/{artifact.id}.pptx').status_code == 401


def test_artifact_pptx_is_for_decks_only(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    principal, _ = in_app(client, seed_starter)
    document = in_app(
        client,
        lambda: workspace.import_directory(principal.workspace_id, 'Doc', 'document', ROOT / 'examples' / 'document'),
    )
    sign_in(client, principal)
    stub_chrome(monkeypatch, 200, b'', json.dumps(SCENE))
    response = client.get(f'/artifacts/{document.id}.pptx')
    assert response.status_code == 422 and 'only a deck' in response.json()['detail']
    # And without a chrome service a deck's export is a 503 like the PDF.
    monkeypatch.delenv('OPENARTIFACT_CHROME_URL')
    artifact = starter(client)
    assert client.get(f'/artifacts/{artifact.id}.pptx').status_code == 503


OUTLINE = {
    'kind': 'document',
    'pages': [
        {
            'index': 1,
            'blocks': [
                {
                    'type': 'heading',
                    'level': 1,
                    'runs': [
                        {
                            'text': 'A document',
                            'bold': False,
                            'italic': False,
                            'code': False,
                            'underline': False,
                            'strike': False,
                        }
                    ],
                }
            ],
        }
    ],
}


def test_artifact_docx_is_assembled_from_the_chrome_services_outline(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
):
    principal, _ = in_app(client, seed_starter)
    document = in_app(
        client,
        lambda: workspace.import_directory(principal.workspace_id, 'Doc', 'document', ROOT / 'examples' / 'document'),
    )
    sign_in(client, principal)
    asked = stub_chrome(monkeypatch, 200, b'', json.dumps(OUTLINE))
    response = client.get(f'/artifacts/{document.id}.docx')
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == word.MEDIA_TYPE
    sha = last_commit(client, document)
    assert response.headers['content-disposition'].startswith(
        f'attachment; filename="OpenArtifact Document {sha}.docx"'
    )
    # The outline was read from the print pass in scene mode; nothing was photographed.
    [outline_url] = asked
    assert re.fullmatch(rf'http://app:8765/print/[^/]+/artifacts/{document.id}/\?scene', outline_url)
    [heading] = Document(io.BytesIO(response.content)).paragraphs
    assert (heading.style.name, heading.text) == ('Heading 1', 'A document')  # pyright: ignore
    # A page artifact exports the same way; a deck does not, and a visitor is sent away.
    page = in_app(
        client, lambda: workspace.import_directory(principal.workspace_id, 'Page', 'page', ROOT / 'examples' / 'page')
    )
    asked = stub_chrome(monkeypatch, 200, b'', json.dumps({**OUTLINE, 'kind': 'page'}))
    response = client.get(f'/artifacts/{page.id}.docx')
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == word.MEDIA_TYPE
    [outline_url] = asked
    assert re.fullmatch(rf'http://app:8765/print/[^/]+/artifacts/{page.id}/\?scene', outline_url)
    deck = starter(client)
    response = client.get(f'/artifacts/{deck.id}.docx')
    assert response.status_code == 422 and 'only a document or page' in response.json()['detail']
    sign_out(client)
    assert client.get(f'/artifacts/{document.id}.docx').status_code == 401


def test_exports_of_unknown_artifacts_are_404(client: TestClient):
    missing = uuid.uuid4()
    for suffix in ('.md', '.zip', '.pdf', '.png', '.pptx', '.docx'):
        assert client.get(f'/artifacts/{missing}{suffix}').status_code == 404, suffix
        assert client.get(f'/artifacts/not-a-uuid{suffix}').status_code == 404, suffix


FAR_FUTURE = 4_000_000_000


def upload_to(artifact: workspace.Artifact, path: str, size: int, expires: int = FAR_FUTURE) -> str:
    """The URL the `upload_url` tool would mint, built directly so these tests need no MCP call."""
    return f'/artifacts/{artifact.id}/{path}?token={upload.make_token(artifact.id, path, size, expires)}'


def test_upload_writes_and_commits(client: TestClient):
    artifact = starter(client)
    body = b'\x89PNG not really'
    response = client.put(upload_to(artifact, 'assets/new/pic.png', len(body)), content=body)
    assert response.status_code == 200, response.text
    assert response.json() == {
        'path': 'assets/new/pic.png',
        'size': len(body),
        'sha256': hashlib.sha256(body).hexdigest(),
    }
    # The file is in the checkout, committed on its own, and served back.
    checkout = workspace.checkout_path(artifact.id)
    assert (checkout / 'assets' / 'new' / 'pic.png').read_bytes() == body
    log = in_app(client, lambda: workspace.artifact_git(artifact.id, 'log', '--format=%s')).splitlines()
    assert log[0] == 'upload: assets/new/pic.png'
    assert client.get(f'/artifacts/{artifact.id}/assets/new/pic.png').content == body
    # Overwriting an existing file works the same way.
    again = client.put(upload_to(artifact, 'main.md', 6), content=b'# new\n')
    assert again.status_code == 200
    assert client.get(f'/artifacts/{artifact.id}/main.md').text == '# new\n'


def test_upload_rejects_bad_tokens(client: TestClient):
    artifact = starter(client)
    # Wrong size, wrong path, wrong artifact, expired, garbage, missing.
    url = upload_to(artifact, 'assets/x.png', 3)
    assert client.put(url, content=b'four').status_code == 403
    token = url.partition('?token=')[2]
    assert client.put(f'/artifacts/{artifact.id}/assets/y.png?token={token}', content=b'abc').status_code == 403
    other = starter(client)
    assert client.put(f'/artifacts/{other.id}/assets/x.png?token={token}', content=b'abc').status_code == 403
    expired = client.put(upload_to(artifact, 'assets/x.png', 3, expires=1_000), content=b'abc')
    assert expired.status_code == 403 and 'expired' in expired.json()['detail']
    assert client.put(f'/artifacts/{artifact.id}/assets/x.png?token=junk', content=b'abc').status_code == 403
    assert client.put(f'/artifacts/{artifact.id}/assets/x.png', content=b'abc').status_code == 403
    # Nothing was written or committed.
    checkout = workspace.checkout_path(artifact.id)
    assert not (checkout / 'assets' / 'x.png').exists()
    assert not (checkout / 'assets' / 'y.png').exists()
    log = in_app(client, lambda: workspace.artifact_git(artifact.id, 'log', '--format=%s')).splitlines()
    assert not any(line.startswith('upload:') for line in log)


def test_upload_path_is_checked_even_with_a_valid_token(client: TestClient):
    """A token for a bad path cannot be minted by the tool; the route refuses one anyway."""
    artifact = starter(client)
    for path in ('dist/index.html', 'assets/.gitignore'):
        response = client.put(upload_to(artifact, path, 3), content=b'abc')
        assert response.status_code == 403, path
    # A directory in the way is a conflict, not a crash.
    response = client.put(upload_to(artifact, 'components', 3), content=b'abc')
    assert response.status_code == 409
    response = client.put(upload_to(artifact, 'main.md/x.txt', 3), content=b'abc')
    assert response.status_code == 409


def test_upload_cannot_follow_a_symlink_out(client: TestClient, tmp_path: Path):
    artifact = starter(client)
    outside = tmp_path / 'outside'
    outside.mkdir()
    directory = workspace.checkout_path(artifact.id)
    (directory / 'assets' / 'link').symlink_to(outside)
    response = client.put(upload_to(artifact, 'assets/link/x.png', 3), content=b'abc')
    assert response.status_code == 403
    assert not (outside / 'x.png').exists()


def test_upload_needs_a_content_length(client: TestClient):
    artifact = starter(client)

    def chunks() -> Iterator[bytes]:
        yield b'abc'

    # A generator body is sent chunked, without Content-Length.
    response = client.put(upload_to(artifact, 'assets/x.png', 3), content=chunks())
    assert response.status_code == 411


def test_upload_to_unknown_artifact_is_404(client: TestClient):
    missing = uuid.uuid4()
    token = upload.make_token(missing, 'main.md', 3, FAR_FUTURE)
    assert client.put(f'/artifacts/{missing}/main.md?token={token}', content=b'abc').status_code == 404
    assert client.put('/artifacts/not-a-uuid/main.md?token=x', content=b'abc').status_code == 404


def test_mcp_without_trailing_slash_is_served_in_place(client: TestClient):
    """`/mcp` answers as `/mcp/` does, with no redirect: claude.ai does not follow one."""
    response = client.post(
        '/mcp', headers={'accept': 'application/json, text/event-stream'}, json={}, follow_redirects=False
    )
    assert response.status_code == 401
    assert response.headers['www-authenticate'] == client.post('/mcp/', json={}).headers['www-authenticate']
    metadata = client.get('/.well-known/oauth-protected-resource/mcp', follow_redirects=False)
    assert metadata.status_code == client.get('/.well-known/oauth-protected-resource/mcp/').status_code != 307


def test_mcp_requires_a_token(client: TestClient):
    response = client.post(
        '/mcp/',
        headers={'accept': 'application/json, text/event-stream'},
        json={
            'jsonrpc': '2.0',
            'id': 1,
            'method': 'initialize',
            'params': {
                'protocolVersion': '2025-06-18',
                'capabilities': {},
                'clientInfo': {'name': 't', 'version': '0'},
            },
        },
    )
    assert response.status_code == 401


@pytest.fixture
def live_server(server_env: None) -> Iterator[str]:
    """Run the app in a uvicorn thread on a free port; the MCP transport needs a real HTTP connection."""
    config = uvicorn.Config(server.app, host='127.0.0.1', port=0, log_level='warning')
    live = uvicorn.Server(config)
    thread = threading.Thread(target=live.run, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10
    while not live.started:
        # A failing lifespan ends the thread without ever setting `started`; do not spin forever on it.
        if not thread.is_alive() or time.monotonic() > deadline:
            raise RuntimeError('uvicorn did not start')
        thread.join(0.01)
    port = live.servers[0].sockets[0].getsockname()[1]
    try:
        yield f'http://127.0.0.1:{port}'
    finally:
        live.should_exit = True
        thread.join(5)


@pytest.mark.anyio
async def test_mcp_over_http(live_server: str):
    async with Client(f'{live_server}/mcp/', auth=DEV_TOKEN) as client:
        tools = {tool.name for tool in await client.list_tools()}
        assert tools == {
            'new_personal_artifact',
            'new_org_artifact',
            'run_code',
            'build',
            'screenshot',
            'upload_url',
            'set_access',
            'fork',
            'list_artifacts',
        }
        created = await client.call_tool(
            'new_personal_artifact', {'title': 'Demo', 'content': '# Demo\n', 'public': True}
        )
        name = created.data.partition('\n')[0].removeprefix('artifact: ')
        result = await client.call_tool(
            'run_code',
            {'artifact': name, 'code': "from pathlib import Path\nPath('main.md').write_text('# Demo\\n')"},
        )
        assert result.data == '7\n'
        built = await client.call_tool('build', {'artifact': name})
        assert built.data.endswith(f'page: http://127.0.0.1:8765/artifacts/{name}/\n')
        listed = await client.call_tool('list_artifacts', {})
        assert listed.data.startswith(f'{name}  deck  Demo  [public]  ')
        logo = (STARTER / 'assets' / 'logo.svg').read_bytes()
        minted = await client.call_tool('upload_url', {'artifact': name, 'files': [['assets/logo.svg', len(logo)]]})
        [upload_url] = minted.data
    async with httpx2.AsyncClient() as http:
        # The page the tool pointed at is served by the same process.
        page = await http.get(f'{live_server}/artifacts/{name}/')
        assert page.status_code == 200
        assert 'id="artifact-markdown"' in page.text
        # The minted URL carries the configured base URL; this server is on a free port.
        uploaded = await http.put(upload_url.replace('http://127.0.0.1:8765', live_server), content=logo)
        assert uploaded.status_code == 200, uploaded.text
        assert uploaded.json()['sha256'] == hashlib.sha256(logo).hexdigest()
        served = await http.get(f'{live_server}/artifacts/{name}/assets/logo.svg')
        assert served.content == logo


@pytest.mark.anyio
async def test_native_telemetry_reaches_logfire(live_server: str, capfire: CaptureLogfire):
    """FastAPI, FastMCP, asyncpg and monty spans all land in Logfire with no `instrument_*` beyond the two hooks."""
    logfire.instrument_asyncpg()
    logfire.instrument_monty()
    async with httpx2.AsyncClient() as http:
        assert (await http.get(f'{live_server}/api/')).status_code == 200
    async with Client(f'{live_server}/mcp/', auth=DEV_TOKEN) as client:
        created = await client.call_tool('new_personal_artifact', {'title': 'T', 'content': '# T\n', 'public': False})
        name = created.data.partition('\n')[0].removeprefix('artifact: ')
        await client.call_tool('run_code', {'artifact': name, 'code': "print('hi')"})
    spans = capfire.exporter.exported_spans_as_dict()
    names = {span['name'] for span in spans}
    # Exported names are message templates: FastAPI's route, FastMCP's tool, monty's session and run, our git spans.
    assert {
        'GET /api/',
        'fastapi.endpoint',
        # The MCP app is a root mount, so FastAPI's telemetry knows the request only as the mount's path.
        'POST /{path}',
        'tools/call run_code',
        'session {script_name}',
        'run code',
        'git {argv}',
    } <= names
    assert any(span['attributes'].get('db.system') == 'postgresql' for span in spans)


# --- access ------------------------------------------------------------------


def set_access(client: TestClient, artifact: workspace.Artifact, visibility: str, org_editable: bool = False) -> None:
    """Change the permissions, keeping the placement the artifact has."""

    async def change() -> workspace.Artifact:
        current = await workspace.get_artifact(artifact.id)
        assert current is not None
        return await workspace.set_access(
            artifact.id, visibility=visibility, org_editable=org_editable, organization_id=current.organization_id
        )

    in_app(client, change)


def org_members(client: TestClient) -> tuple[auth.Principal, auth.Principal]:
    """A colleague of the seed user (same Workspace domain) and an outsider, created through the app."""

    async def make() -> tuple[auth.Principal, auth.Principal]:
        async with db.pool().acquire() as conn, conn.transaction():
            await auth.upsert_user(conn, sub='seed', email='seed@x.test', name='Seed', picture=None, hd='x.test')
            colleague = await auth.upsert_user(conn, sub='c', email='c@x.test', name='C', picture=None, hd='x.test')
            outsider = await auth.upsert_user(conn, sub='s', email='s@gmail.test', name='S', picture=None)
        return colleague, outsider

    return in_app(client, make)


def org_artifact(client: TestClient, artifact: workspace.Artifact, colleague: auth.Principal) -> None:
    """Move the seed artifact into the organisation (the row's org, as `new_org_artifact` would set it)."""
    org_id = colleague.organization_id
    assert org_id is not None
    in_app(
        client,
        lambda: db.pool().execute(
            "UPDATE artifacts SET organization_id = $2, visibility = 'org' WHERE id = $1", artifact.id, org_id
        ),
    )


ROUTES = (
    '/artifacts/{id}/',
    '/artifacts/{id}/main.md',
    '/artifacts/{id}/assets/logo.svg',
    '/artifacts/{id}.md',
    '/artifacts/{id}.zip',
    '/artifacts/{id}.json',
)


def statuses(client: TestClient, artifact: workspace.Artifact) -> set[int]:
    """The status every artifact route answers with, as a JSON client (no redirects)."""
    headers = {'accept': 'application/json'}
    return {client.get(route.format(id=artifact.id), headers=headers).status_code for route in ROUTES}


def test_private_artifact_is_the_owners_alone(client: TestClient):
    artifact = starter(client)
    colleague, outsider = org_members(client)
    assert statuses(client, artifact) == {200}
    # A visitor: a browser is sent to sign in, anything else gets a 401 that says where.
    sign_out(client)
    assert statuses(client, artifact) == {401}
    page = client.get(f'/artifacts/{artifact.id}/', headers={'accept': 'text/html'}, follow_redirects=False)
    assert page.status_code == 303
    assert page.headers['location'] == f'/login?next=%2Fartifacts%2F{artifact.id}%2F'
    denied = client.get(f'/artifacts/{artifact.id}.json')
    assert denied.status_code == 401 and denied.json()['login_url'] == f'/login?next=%2Fartifacts%2F{artifact.id}%2F'
    # Signed in but not shared with: a 403 page for a browser, JSON otherwise. Colleagues too: it is personal.
    for who in (colleague, outsider):
        sign_in(client, who)
        assert statuses(client, artifact) == {403}
        page = client.get(f'/artifacts/{artifact.id}/', headers={'accept': 'text/html'})
        assert page.status_code == 403 and 'This artifact is private' in page.text and (who.email or '') in page.text
        assert page.headers['content-security-policy'].startswith("default-src 'none'")
    # The upload route is governed by its own token, not the session.
    sign_out(client)
    assert client.put(upload_to(artifact, 'x.txt', 1), content=b'x').status_code == 200


def test_org_visible_and_editable(client: TestClient):
    artifact = starter(client)
    colleague, outsider = org_members(client)
    org_artifact(client, artifact, colleague)
    sign_in(client, colleague)
    assert statuses(client, artifact) == {200}
    info = client.get(f'/artifacts/{artifact.id}.json').json()
    assert info['visibility'] == 'org' and info['organization'] == {'domain': 'x.test', 'name': 'x.test'}
    assert info['viewer'] == {'name': 'C', 'email': 'c@x.test', 'picture': None}
    assert (info['can_edit'], info['can_fork'], info['org_editable']) == (False, True, False)
    sign_in(client, outsider)
    assert statuses(client, artifact) == {403}
    sign_out(client)
    assert statuses(client, artifact) == {401}
    set_access(client, artifact, 'org', org_editable=True)
    sign_in(client, colleague)
    assert client.get(f'/artifacts/{artifact.id}.json').json()['can_edit'] is True


def test_public_artifact_is_open_to_all(client: TestClient):
    artifact = starter(client)
    set_access(client, artifact, 'public')
    sign_out(client)
    assert statuses(client, artifact) == {200}
    info = client.get(f'/artifacts/{artifact.id}.json').json()
    assert info['viewer'] is None and info['can_fork'] is False and info['can_edit'] is False
    assert info['login_url'] == f'/login?next=%2Fartifacts%2F{artifact.id}%2F'
    page = client.get(f'/artifacts/{artifact.id}/')
    assert page.headers['cache-control'] == 'private' and page.headers['vary'] == 'Accept, Cookie'


def test_browser_fork(client: TestClient):
    artifact = starter(client)
    colleague, _outsider = org_members(client)
    org_artifact(client, artifact, colleague)
    # Anonymous: sign in first. Cross-site: refused.
    sign_out(client)
    assert client.post(f'/artifacts/{artifact.id}/fork').status_code == 401
    sign_in(client, colleague)
    assert client.post(f'/artifacts/{artifact.id}/fork', headers={'sec-fetch-site': 'cross-site'}).status_code == 403
    response = client.post(f'/artifacts/{artifact.id}/fork', follow_redirects=False)
    assert response.status_code == 303
    location = response.headers['location']
    fork_id = uuid.UUID(location.removeprefix('/artifacts/').rstrip('/'))
    fork = in_app(client, lambda: workspace.get_artifact(fork_id))
    assert fork is not None and fork.forked_from == artifact.id and fork.workspace_id == colleague.workspace_id
    assert (fork.visibility, fork.organization_id) == ('private', None)
    info = client.get(f'{location[:-1]}.json').json()
    assert info['forked_from'] == str(artifact.id) and info['can_edit'] is True
    assert client.get(f'{location}main.md').content == (STARTER / 'main.md').read_bytes()
    log = in_app(client, lambda: workspace.artifact_git(fork_id, 'log', '--format=%s')).splitlines()
    assert log[0] == f'fork of {artifact.id}'


@pytest.mark.filterwarnings('ignore:A configured store is unstable')
def test_oauth_metadata_points_at_served_routes(monkeypatch: pytest.MonkeyPatch):
    """With Google configured, every URL the OAuth metadata advertises is a route the app serves.

    The MCP app is mounted at the root with its endpoint at `/mcp/`, which is what makes this hold: mounted under
    `/mcp`, the routes would sit one level below the URLs built from the base URL.
    """
    from cryptography.fernet import Fernet
    from fastmcp import FastMCP

    monkeypatch.setenv('GOOGLE_CLIENT_ID', 'client-id.apps.googleusercontent.com')
    monkeypatch.setenv('GOOGLE_CLIENT_SECRET', 'client-secret')
    monkeypatch.setenv('OPENARTIFACT_SECRET_KEY', Fernet.generate_key().decode())
    monkeypatch.setenv('OPENARTIFACT_BASE_URL', 'https://example.com')
    provider = auth.make_auth_provider()
    throwaway = FastAPI()
    throwaway.mount('/', FastMCP('t', auth=provider).http_app(path='/mcp/'))
    with TestClient(throwaway, base_url='https://example.com') as http:
        metadata = http.get('/.well-known/oauth-authorization-server').json()
        urls = {metadata[key] for key in ('authorization_endpoint', 'token_endpoint', 'registration_endpoint')}
        assert urls == {f'https://example.com/{p}' for p in ('authorize', 'token', 'register')}
        for url in urls:
            # GET on a POST-only route is 405, which still proves the route is there; a miss would be 404.
            assert http.get(url).status_code != 404, url
        challenge = http.post('/mcp/', json={}).headers['www-authenticate']
        resource = re.search(r'resource_metadata="([^"]+)"', challenge)
        assert resource and http.get(resource.group(1)).status_code == 200, challenge
    # Building the provider installed our consent and error pages in FastMCP's modules.
    import fastmcp.server.auth.oauth_proxy.consent as fastmcp_consent
    import fastmcp.server.auth.oauth_proxy.proxy as fastmcp_proxy

    import pages

    assert vars(fastmcp_consent)['create_consent_html'] is pages.consent_html
    assert vars(fastmcp_proxy)['create_error_html'] is pages.oauth_error_html


def test_our_consent_page_matches_fastmcp_contract():
    """`pages.consent_html` accepts every argument FastMCP's renderer takes, and renders the form it expects."""
    import inspect

    from fastmcp.server.auth.oauth_proxy import ui

    import pages

    theirs = inspect.signature(ui.create_consent_html).parameters
    ours = inspect.signature(pages.consent_html).parameters
    accepts_rest = any(p.kind is inspect.Parameter.VAR_KEYWORD for p in ours.values())
    named = {name for name, p in ours.items() if p.kind is inspect.Parameter.KEYWORD_ONLY}
    assert accepts_rest and {'client_id', 'redirect_uri', 'scopes', 'txn_id', 'csrf_token'} <= named <= set(theirs)
    html = pages.consent_html(
        client_id='c',
        redirect_uri='https://x/cb',
        scopes=['openid'],
        txn_id='t1',
        csrf_token='k1',
        client_name='Claude',
    )
    for needle in (
        'name="txn_id" value="t1"',
        'name="csrf_token" value="k1"',
        'value="approve"',
        'value="deny"',
        'Claude',
    ):
        assert needle in html
    theirs_error = inspect.signature(ui.create_error_html).parameters
    assert {'error_title', 'error_message', 'error_details'} <= set(theirs_error)
    assert '<h1>Oops</h1>' in pages.oauth_error_html('Oops', 'it broke', {'code': 'x'})
