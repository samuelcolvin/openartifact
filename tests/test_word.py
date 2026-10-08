"""Tests for `backend/word.py`: the outline to Word mapping, read back with python-docx. No Postgres, no Chrome."""

# python-docx types little of what is read back.
# pyright: basic

from __future__ import annotations

import io
from pathlib import Path

import word
from docx import Document
from docx.oxml.ns import qn
from PIL import Image


def run(text: str, **overrides: object) -> dict[str, object]:
    base: dict[str, object] = {
        'text': text,
        'bold': False,
        'italic': False,
        'code': False,
        'underline': False,
        'strike': False,
    }
    return {**base, **overrides}


OUTLINE = {
    'kind': 'document',
    'pages': [
        {
            'index': 1,
            'blocks': [
                {'type': 'heading', 'level': 1, 'runs': [run('Title')]},
                {
                    'type': 'paragraph',
                    'runs': [
                        run('Plain, '),
                        run('bold', bold=True),
                        run(' and '),
                        run('code', code=True),
                        run('\n'),
                        run('link', href='https://example.com/'),
                    ],
                },
                {
                    'type': 'list',
                    'ordered': True,
                    'items': [
                        {
                            'runs': [run('one')],
                            'blocks': [
                                {'type': 'list', 'ordered': False, 'items': [{'runs': [run('nested')], 'blocks': []}]}
                            ],
                        },
                        {'runs': [run('two')], 'blocks': []},
                    ],
                },
                {'type': 'list', 'ordered': True, 'items': [{'runs': [run('again from one')], 'blocks': []}]},
                {'type': 'code', 'text': 'def f():\n    return 1'},
                {'type': 'quote', 'blocks': [{'type': 'paragraph', 'runs': [run('quoted')]}]},
                {
                    'type': 'table',
                    'rows': [[{'runs': [run('h')], 'header': True}], [{'runs': [run('d')], 'header': False}]],
                },
                {'type': 'image', 'src': 'assets/pic.png', 'alt': 'A picture'},
                {'type': 'image', 'src': 'assets/logo.svg', 'alt': 'A logo'},
                {'type': 'rule'},
            ],
        },
        {'index': 2, 'blocks': [{'type': 'paragraph', 'runs': [run('second page')]}]},
    ],
}


def test_build_docx_writes_word_paragraphs(tmp_path: Path):
    outline = word.Outline.model_validate(OUTLINE)
    (tmp_path / 'assets').mkdir()
    Image.new('RGB', (4, 2)).save(tmp_path / 'assets' / 'pic.png', 'PNG')
    (tmp_path / 'assets' / 'logo.svg').write_text('<svg/>')
    (tmp_path / 'assets' / 'outside.png').symlink_to(Path('/etc/hosts'))
    images = word.load_images(outline, tmp_path)
    assert list(images) == ['assets/pic.png']  # the SVG is not a raster file

    document = Document(io.BytesIO(word.build_docx(outline, images, 'Doc')))
    assert document.core_properties.title == 'Doc'
    paragraphs = document.paragraphs
    styles = [(p.style.name, p.text) for p in paragraphs]  # pyright: ignore
    assert styles[0] == ('Heading 1', 'Title')
    assert styles[1] == ('Normal', 'Plain, bold and code\nlink')
    assert styles[2:6] == [
        ('List Number', 'one'),
        ('List Bullet 2', 'nested'),
        ('List Number', 'two'),
        ('List Number', 'again from one'),
    ]
    assert styles[6] == ('No Spacing', 'def f():\n    return 1')
    assert styles[7] == ('Quote', 'quoted')
    assert ('Normal', '[image: A logo]') in styles
    assert styles[-1] == ('Normal', 'second page')

    # Run styles: a bold run, a monospace code run, and the link as a hyperlink with its own run.
    body = paragraphs[1]
    bold, code = body.runs[1], body.runs[3]
    assert bold.text == 'bold' and bold.font.bold and code.text == 'code' and code.font.name == 'Consolas'
    link = body._p.find(qn('w:hyperlink'))
    assert link is not None and link.find(f'{qn("w:r")}/{qn("w:t")}').text == 'link'
    assert document.part.rels[link.get(qn('r:id'))].target_ref == 'https://example.com/'
    # No colour anywhere: Word's defaults throughout.
    assert all(r.font.color.rgb is None for p in paragraphs for r in p.runs)

    # The second ordered list restarts: it has a numbering instance of its own.
    first, second = paragraphs[2], paragraphs[5]
    assert first._p.pPr.numPr.numId.val != second._p.pPr.numPr.numId.val  # pyright: ignore
    assert paragraphs[3]._p.pPr.numPr is None  # pyright: ignore

    # A table with a bold header row, one picture, a rule as a bottom border, and a page break before page two.
    [table] = document.tables
    assert table.style.name == 'Table Grid' and table.rows[0].cells[0].paragraphs[0].runs[0].font.bold  # pyright: ignore
    assert len(document.inline_shapes) == 1
    assert any(p._p.pPr is not None and p._p.pPr.find(qn('w:pBdr')) is not None for p in paragraphs)
    assert document.element.body.findall(f'.//{qn("w:br")}[@{qn("w:type")}="page"]')


def test_walk_finds_images_in_lists_and_quotes():
    outline = word.Outline.model_validate(OUTLINE)
    assert [image.src for image in word.walk(outline.pages[0].blocks)] == ['assets/pic.png', 'assets/logo.svg']
    nested = word.Outline.model_validate(
        {
            'kind': 'document',
            'pages': [
                {
                    'index': 1,
                    'blocks': [
                        {'type': 'quote', 'blocks': [{'type': 'image', 'src': 'a.png', 'alt': ''}]},
                        {
                            'type': 'list',
                            'ordered': False,
                            'items': [{'runs': [], 'blocks': [{'type': 'image', 'src': 'b.png', 'alt': ''}]}],
                        },
                    ],
                }
            ],
        }
    )
    assert [image.src for image in word.walk(nested.pages[0].blocks)] == ['a.png', 'b.png']
